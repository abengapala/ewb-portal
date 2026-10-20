"""EWB Field Result Autostat Tab - for use inside the EWB Portal multi-page app.

All logic is preserved from the standalone app.py exactly.
Only the top-level st.set_page_config / st.title / __main__ guard are removed
so this can be imported safely as a tab inside app.py.
"""
import csv
import io
import re
import zipfile
from datetime import datetime, time
from pathlib import Path

import msoffcrypto
import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

DEFAULT_PASSWORD = "spmadrid1234"
OPEN_PASSWORDS = [DEFAULT_PASSWORD, "SPM2026", "SPM 2026"]
DATA_DIR = Path(__file__).resolve().parent.parent / "ewb_data"
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

AUTO_BANK = "EWB AUTO LOAN RECOV"
AUTO_REMARKS_BY = "SRRUFO"
AUTO_TIME = "08:30:00"
STATUS_MAP_FILE = DATA_DIR / "status_map.csv"
ACCT_MAP_FILE   = DATA_DIR / "chcode_accounts.csv"
DG_FILE         = DATA_DIR / "datagrid_accounts.csv"
AUTO_COLS = ["CH CODE", "ACCOUNT NUMBER", "STATUS CODE", "REMARKS", "REMARKS BY", "REMARKS DATE"]
FVR_FIELDS = ["chcode", "CH NAME", "status", "SUB STATUS", "EMAIL ADDRESS", "ADDTYPE", "INFORMANT",
              "3RD PARTY LIST", "CLIENT NUMBER", "Message", "field_name",
              "DATE", "TIME", "BANK", "ACCOUNT NUMBER", "reference_code"]
REMARK_FIELDS = ["status", "SUB STATUS", "EMAIL ADDRESS", "CLIENT NUMBER", "ADDTYPE", "INFORMANT",
                 "3RD PARTY LIST", "Message", "field_name"]
BLANKS = {"", "0", "0.0", "none", "nan", "null"}

REF_DEFAULT = """REPO-PAYMENT|FIELD - KEPT_FULL UPDATE
PTP-PTP - PAYMENT|FIELD - PTP_FULL UPDATE
PTP-PTP - REPO|FIELD - PTP REPO
REPO-REPO|FIELD - KEPT_REPO CLIENT
PTP-PTP - FOLLOW UP VISIT|FIELD - FOLLOW UP COMPLIANT
TP-UNDERNEGO|FIELD - FOLLOW UP COMPLIANT
TP-LEAVE A MESSAGE|FIELD - LEAVE_MESSAGE TO 3RD PARTY
TP-UNIT UNDER INSURANCE|FIELD - INSURANCE CLAIM
TP-NO INTENTION TO PAY OR SURRENDER|FIELD - NO INTENTION TO PAY
TP-CARNAPPED|FIELD - UNIT CARNAPPED
TP-TOTAL WRECKED OR UNIT DAMAGED|FIELD - UNIT DAMAGE OR WRECK
TP-NO CAPACITY TO PAY|FIELD - NO INTENTION TO PAY
TP-UNIT CARNAPPED|FIELD - UNIT CARNAPPED
POS W/C-CLIENT OUT OF AREA|FIELD - CLIENT_OUT OF AREA
POS W/C-CLIENT REFUSED TO TALK|FIELD - NO INTENTION TO PAY
POS W/C-CLIENT DECEASED|FIELD - DECEASED
POS-CLIENT OUT OF AREA|FIELD - CLIENT_OUT OF AREA
POS-NOT ALLOWED TO ENTER - VERIFIED ADDRESS|FIELD - NOT_ALLOWED TO ENTER
POS-HOUSE CLOSED - VERIFIED|FIELD - HOUSED CLOSED VERIFIED
POS-CLIENT REFUSED TO TALK|FIELD - NO INTENTION TO PAY
POS-OOTC (OUT OF THE COUNTRY)|FIELD - CLIENT_OUT OF AREA
POS-CLIENT OUT OF THE COUNTRY|FIELD - CLIENT_OUT OF AREA
POS-NEW HOUSE ADDRESS|FIELD - FOLLOW UP COMPLIANT
POS-CLIENT OUT OF OFFICE|FIELD - CLIENT_OUT OF AREA
POS-IS AT WORK|FIELD - CLIENT_OUT OF AREA
POS-UNOCCUPIED - VERIFIED ADDRESS|FIELD - HOUSED CLOSED VERIFIED
POS-LOT ONLY - VERIFIED ADDRESS|FIELD - LOT_ONLY
POS-REFUSED TO TALK|FIELD - NO INTENTION TO PAY
POS-HOUSE CLOSED|FIELD - HOUSED CLOSED VERIFIED
POS-OUT OF AREA|FIELD - CLIENT_OUT OF AREA
POS-IS ON LEAVE|FIELD - CLIENT_OUT OF AREA
POS-HOUSE CLOSED - STILL RESIDING|FIELD - HOUSED CLOSED VERIFIED
POS-UNDERNEGO|FIELD - FOLLOW UP COMPLIANT
NEG-HOUSE CLOSED - UNVERIFIED|FIELD - HOUSED_CLOSED UNVERIFIED
NEG-CLIENT UNKNOWN|FIELD - CLIENT_UNKNOWN
NEG-NOT ALLOWED TO ENTER - UNVERIFIED ADDRESS|FIELD - NOT_ALLOWED TO ENTER
NEG-ADDRESS UNLOCATED|FIELD - UNLOCATED
NEG-CLIENT MOVED OUT|FIELD - MOVED_OUT
NEG-LOT ONLY - UNVERIFIED ADDRESS|FIELD - LOT_ONLY
NEG-UNOCCUPIED - UNVERIFIED ADDRESS|FIELD - HOUSED_CLOSED UNVERIFIED
NEG-CLIENT RESIGNED|FIELD - UNLOCATED
NEG-ABANDONED|FIELD - MOVED_OUT
NEG-DECEASED|FIELD - DECEASED
NEG-UNLOCATED|FIELD - UNLOCATED
CANCEL-PAID OR REPOSSESSED|FIELD - RESULT
CANCEL-PULLED OUT|FIELD - RESULT
CANCEL-HOLD WITH ARRANGEMENT|FIELD - RESULT
CANCEL-HOLD BY THE BANK|FIELD - RESULT
CANCEL-CANCEL BY AGENT|FIELD - RESULT
CANCEL-INCOMPLETE ADDRESS|FIELD - RESULT
CANCEL-CRITICAL AREA|FIELD - RESULT
CANCEL-PULL OUT BY THE BANK|FIELD - RESULT
TP-CLAIMING FULLY PAID|FIELD - CLAIMING PAID
TP-UNIT IMPOUNDED|FIELD - UNIT IMPOUNDED
POS W/C-HOUSE CLOSED - VERIFIED|FIELD - HOUSED CLOSED VERIFIED
POS W/C-LOT ONLY - VERIFIED ADDRESS|FIELD - LOT_ONLY
POS W/C-CLIENT OUT OF THE COUNTRY|FIELD - CLIENT_OUT OF AREA
POS-HOUSE ABANDONED|FIELD - MOVED_OUT
POS-OUT OF OFFICE|FIELD - CLIENT_OUT OF AREA
POS-UNOCCUPIED|FIELD - MOVED_OUT
NEG-MOVE OUT|FIELD - MOVED_OUT
NEG-UNKNOWN|FIELD - CLIENT_UNKNOWN
TP-UNIT WITH ASSUMER|FIELD - UNIT ASSUMED
POS-CLIENT DECEASED|FIELD - DECEASED
POS-NOT AROUND|FIELD - CLIENT_OUT OF AREA
POS-OOT (OUT OF TOWN)|FIELD - CLIENT_OUT OF AREA
POS W/C-NEW HOUSE ADDRESS|FIELD - FOLLOW UP COMPLIANT
POS W/C-NOT ALLOWED TO ENTER - VERIFIED ADDRESS|FIELD - NOT_ALLOWED TO ENTER
POS W/C-CLIENT OUT OF OFFICE|FIELD - CLIENT_OUT OF AREA
NEG-TRANSFERRED TO OTHER BRANCH|FIELD - CLIENT_OUT OF AREA
TP-SCAMMED|FIELD - UNIT CARNAPPED
POS-CLAIMING FULLY PAID|FIELD - CLAIMING PAID
RESOLVED-REPO|FIELD - KEPT_REPO CLIENT
RESOLVED-PAYMENT|FIELD - KEPT_FULL UPDATE
NEG-DENIED ENTRY|FIELD - NOT_ALLOWED TO ENTER"""

STATUS_ALIGN = {
    "POS-OUT OF AREA-CLIENT NOT AROUND": "FIELD - CLIENT_OUT OF AREA",
    "POS W/C-OUT OF AREA-CLIENT NOT AROUND": "FIELD - CLIENT_OUT OF AREA",
    "POS-CLIENT IS AT WORK": "FIELD - CLIENT_OUT OF AREA",
    "POS W/C-CLIENT IS AT WORK": "FIELD - CLIENT_OUT OF AREA",
    "NEG-UNIT NOT SEEN": "FIELD - UNLOCATED",
}
REF_KEYS = {re.sub(r"\s+", " ", ln.split("|", 1)[0]).strip().upper() for ln in REF_DEFAULT.splitlines()}


# ---------------------------------------------------------------- helpers
def to_date(v, dayfirst=False):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return pd.NaT
    if isinstance(v, (datetime, pd.Timestamp)):
        return pd.Timestamp(v).normalize()
    s = str(v).strip()
    if not s:
        return pd.NaT
    if re.fullmatch(r"\d{5}(\.\d+)?", s):
        return (pd.Timestamp("1899-12-30") + pd.to_timedelta(float(s), unit="D")).normalize()
    if re.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}", s):
        try:
            return pd.Timestamp(s[:10]).normalize()
        except Exception:
            pass
    fmts = ("%d/%m/%Y", "%m/%d/%Y", "%d/%m/%y", "%m/%d/%y") if dayfirst else \
           ("%m/%d/%Y", "%d/%m/%Y", "%m/%d/%y", "%d/%m/%y")
    for f in fmts:
        try:
            return pd.Timestamp(datetime.strptime(s[:10].replace("-", "/"), f)).normalize()
        except ValueError:
            pass
    try:
        return pd.to_datetime(s).normalize()
    except Exception:
        return pd.NaT


def _txt(x):
    if isinstance(x, time):
        return x.strftime("%H:%M:%S")
    if x is None:
        return ""
    if isinstance(x, float) and x.is_integer():
        x = int(x)
    return _ILLEGAL.sub("", str(x).strip())


def _bytes(file):
    if isinstance(file, (str, Path)):
        return Path(file).read_bytes()
    if hasattr(file, "seek"):
        file.seek(0)
    return file.read()


def csv_rows(data):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    return list(csv.reader(io.StringIO(text)))


def open_file_wb(file, password=DEFAULT_PASSWORD):
    data = _bytes(file)
    bio = io.BytesIO(data)
    if not zipfile.is_zipfile(bio):
        bio.seek(0)
        ms = msoffcrypto.OfficeFile(bio)
        if not ms.is_encrypted():
            raise ValueError("This is not a valid .xlsx file.")
        last = None
        for pw in dict.fromkeys([password, *OPEN_PASSWORDS]):
            if not pw:
                continue
            try:
                bio.seek(0)
                ms = msoffcrypto.OfficeFile(bio)
                ms.load_key(password=pw)
                dec = io.BytesIO()
                ms.decrypt(dec)
                dec.seek(0)
                return load_workbook(dec, data_only=True, read_only=True)
            except msoffcrypto.exceptions.InvalidKeyError as e:
                last = e
        raise ValueError("Wrong password. Type the correct one in the sidebar.") from last
    bio.seek(0)
    return load_workbook(bio, data_only=True, read_only=True)


def _nk(x):
    return re.sub(r"\s+", " ", str(x)).strip().upper()


def data_key(status, sub):
    parts = [_nk(p) for p in (status, sub)]
    return "-".join(p for p in parts if p and p.lower() not in BLANKS)


def _read_pairs(path):
    if not path.exists():
        return {}
    with open(path, newline="", encoding="utf-8") as f:
        return {r[0]: r[1] for r in csv.reader(f) if len(r) >= 2}


def _write_pairs(path, d):
    DATA_DIR.mkdir(exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(sorted(d.items()))


def load_status_map():
    m = {}
    for line in REF_DEFAULT.splitlines():
        k, v = line.split("|", 1)
        m[_nk(k)] = v.strip()
    m.update({_nk(k): v for k, v in STATUS_ALIGN.items()})
    m.update(_read_pairs(STATUS_MAP_FILE))
    return m


def save_status_map(new):
    cur = _read_pairs(STATUS_MAP_FILE)
    cur.update({_nk(k): v for k, v in new.items() if str(v).strip()})
    _write_pairs(STATUS_MAP_FILE, cur)


def load_acct_map():
    return _read_pairs(ACCT_MAP_FILE)


def save_acct_map(new):
    cur = _read_pairs(ACCT_MAP_FILE)
    cur.update({str(k).strip().upper(): str(v).strip() for k, v in new.items()
                if str(v).strip() and str(v).strip().lower() not in BLANKS})
    _write_pairs(ACCT_MAP_FILE, cur)


def load_dg_map():
    return _read_pairs(DG_FILE)


def read_datagrid(file, password=DEFAULT_PASSWORD):
    wb = open_file_wb(file, password)
    ws = next(iter(wb.worksheets))
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}
    hdr = [str(c or "").strip().lower() for c in rows[0]]
    ic_col = next((i for i, h in enumerate(hdr) if "old" in h and "i" in h and "c" in h), None)
    if ic_col is None:
        ic_col = next((i for i, h in enumerate(hdr) if "new" in h and "i" in h and "c" in h), None)
    ac_col = next((i for i, h in enumerate(hdr) if "account" in h and "no" in h), None)
    if ic_col is None or ac_col is None:
        return {}
    result = {}
    for r in rows[1:]:
        ch = _txt(r[ic_col] if ic_col < len(r) else "").strip().upper()
        ac = _txt(r[ac_col] if ac_col < len(r) else "").strip()
        if ch and ac and ac.lower() not in BLANKS:
            result[ch] = ac
    return result


def read_fvr(file, bank=AUTO_BANK, password=DEFAULT_PASSWORD):
    wb = open_file_wb(file, password)
    result_ws = next((ws for ws in wb.worksheets if ws.title.strip().upper() == "RESULT"), None)
    ws = result_ws or wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return pd.DataFrame(columns=FVR_FIELDS)
    hdr_raw = [str(c or "").strip() for c in rows[0]]
    hdr = [h.lower() for h in hdr_raw]

    def col(name):
        nl = name.lower()
        if nl in hdr:
            return hdr.index(nl)
        for i, h in enumerate(hdr):
            if nl in h:
                return i
        return None

    col_map = {f: col(f) for f in FVR_FIELDS}
    data = []
    for r in rows[1:]:
        def g(f):
            i = col_map.get(f)
            return _txt(r[i] if i is not None and i < len(r) else "")
        bank_val = g("BANK")
        if bank.upper() not in bank_val.upper():
            continue
        row = {f: g(f) for f in FVR_FIELDS}
        data.append(row)
    df = pd.DataFrame(data, columns=FVR_FIELDS) if data else pd.DataFrame(columns=FVR_FIELDS)
    df["DATE"] = df["DATE"].map(lambda v: to_date(v))
    df["chcode"] = df["chcode"].str.strip().str.upper()
    df["ACCOUNT NUMBER"] = df["ACCOUNT NUMBER"].str.strip()
    return df


def build_autostat(fvr, day, status_map, acct_map, remarks_by, use_visit_time, fixed_time,
                   dg_map=None, use_dg=True):
    dg_map = dg_map or {}
    day_ts = pd.Timestamp(day).normalize()
    df = fvr[fvr["DATE"] == day_ts].copy()
    stats = {"rows": len(df), "cancel": 0, "dup": 0, "retimed": 0,
             "ready": 0, "hold_status": 0, "hold_acct": 0, "aligned": {},
             "acct_datagrid": 0, "acct_filled": 0}
    if df.empty:
        return pd.DataFrame(columns=AUTO_COLS), pd.DataFrame(), stats

    # Remove CANCEL
    cancel_mask = df["status"].str.upper().str.startswith("CANCEL")
    stats["cancel"] = int(cancel_mask.sum())
    df = df[~cancel_mask].copy()

    # Remove duplicate reference_codes
    dup_mask = df["reference_code"].str.strip().ne("") & df["reference_code"].str.strip().ne("0")
    dups = df[dup_mask].duplicated(subset=["reference_code"], keep="first")
    stats["dup"] = int(dups.sum())
    df = df[~(dup_mask & dups)].copy()

    # Build KEY
    df["KEY"] = df.apply(lambda r: data_key(r["status"], r["SUB STATUS"]), axis=1)

    # Map STATUS CODE
    df["STATUS CODE"] = df["KEY"].map(lambda k: status_map.get(k.upper(), ""))
    aligned = {}
    for k in df[df["STATUS CODE"] == ""]["KEY"].unique():
        if k.upper() in status_map:
            aligned[k] = status_map[k.upper()]
    stats["aligned"] = aligned

    # Resolve account numbers
    def resolve_acct(row):
        acct = str(row.get("ACCOUNT NUMBER", "")).strip()
        if acct and acct.lower() not in BLANKS:
            return acct, "field"
        ch = str(row.get("chcode", "")).strip().upper()
        if use_dg and ch in dg_map:
            return dg_map[ch], "datagrid"
        if ch in acct_map:
            return acct_map[ch], "filled"
        for _, r2 in fvr[fvr["chcode"] == ch].iterrows():
            a2 = str(r2.get("ACCOUNT NUMBER", "")).strip()
            if a2 and a2.lower() not in BLANKS:
                return a2, "filled"
        return "", "missing"

    acct_results = df.apply(resolve_acct, axis=1, result_type="expand")
    df["ACCOUNT NUMBER"] = acct_results[0]
    df["_acct_src"] = acct_results[1]
    stats["acct_datagrid"] = int((df["_acct_src"] == "datagrid").sum())
    stats["acct_filled"] = int((df["_acct_src"] == "filled").sum())

    # Build REMARKS
    def build_remarks(row):
        parts = []
        for f in REMARK_FIELDS:
            v = str(row.get(f, "")).strip()
            if v and v.lower() not in BLANKS:
                parts.append(v)
        return " | ".join(parts)

    df["REMARKS"] = df.apply(build_remarks, axis=1)
    df["REMARKS BY"] = remarks_by

    # Build REMARKS DATE with +1 minute for multiple same-day visits
    if use_visit_time:
        df["_visit_time"] = df["TIME"].str.strip().replace("", fixed_time).fillna(fixed_time)
    else:
        df["_visit_time"] = fixed_time

    # Retime duplicates
    visit_counts = {}
    remark_dates = []
    for _, row in df.iterrows():
        ch = row["chcode"]
        cnt = visit_counts.get(ch, 0)
        visit_counts[ch] = cnt + 1
        try:
            t = datetime.strptime(str(row["_visit_time"]).strip()[:8], "%H:%M:%S")
        except ValueError:
            try:
                t = datetime.strptime(fixed_time, "%H:%M:%S")
            except ValueError:
                t = datetime.strptime("08:30:00", "%H:%M:%S")
        from datetime import timedelta
        t2 = (datetime.combine(datetime.today(), t.time()) + timedelta(minutes=cnt)).time()
        remark_dates.append(f"{pd.Timestamp(day):%m/%d/%Y} {t2.strftime('%H:%M:%S')}")
        if cnt > 0:
            stats["retimed"] += 1

    df["REMARKS DATE"] = remark_dates

    # Split ready vs hold
    hold_mask = (df["STATUS CODE"] == "") | (df["ACCOUNT NUMBER"].str.strip() == "")
    hold = df[hold_mask].copy()
    ready = df[~hold_mask].copy()
    stats["hold_status"] = int((hold["STATUS CODE"] == "").sum())
    stats["hold_acct"] = int((hold["ACCOUNT NUMBER"].str.strip() == "").sum())
    stats["ready"] = len(ready)

    hold["REASON"] = ""
    hold.loc[hold["STATUS CODE"] == "", "REASON"] = "Status not in REF"
    hold.loc[(hold["STATUS CODE"] != "") & (hold["ACCOUNT NUMBER"].str.strip() == ""), "REASON"] = "Missing account number"

    ready_out = ready[["chcode", "ACCOUNT NUMBER", "STATUS CODE", "REMARKS", "REMARKS BY", "REMARKS DATE"]].copy()
    ready_out.columns = AUTO_COLS
    hold_out = hold[["chcode", "CH NAME", "KEY", "ACCOUNT NUMBER", "STATUS CODE", "REMARKS", "REMARKS DATE", "REASON"]].copy()
    hold_out.columns = ["CH CODE", "CH NAME", "KEY", "ACCOUNT NUMBER", "STATUS CODE", "REMARKS", "REMARKS DATE", "REASON"]
    return ready_out, hold_out, stats


def autostat_workbook(df, password=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "sheet1"
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.fill = hdr_fill
        cell.font = hdr_font
    for _, row in df.iterrows():
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    if not password:
        return buf.getvalue()
    buf.seek(0)
    out = io.BytesIO()
    msoffcrypto.OfficeFile(buf).encrypt(password, out)
    return out.getvalue()


def attention_workbook(df):
    wb = Workbook()
    ws = wb.active
    ws.title = "Needs Attention"
    hdr_fill = PatternFill("solid", fgColor="C00000")
    hdr_font = Font(bold=True, color="FFFFFF")
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.fill = hdr_fill
        cell.font = hdr_font
    for _, row in df.iterrows():
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def dg_load():
    ss = st.session_state
    f = ss.get("auto_dg_tab")
    if f is None:
        return
    try:
        dg = read_datagrid(f, ss.get("file_password", DEFAULT_PASSWORD))
        DATA_DIR.mkdir(exist_ok=True)
        _write_pairs(DG_FILE, dg)
        ss.auto_dg_msg = ("success", f"DataGrid loaded: {len(dg):,} accounts saved.")
    except Exception as e:
        ss.auto_dg_msg = ("error", f"Could not read DataGrid: {e}")


def auto_harvest():
    ss = st.session_state
    files = ss.get("auto_csv_tab") or []
    if not files:
        return
    added = 0
    for f in files:
        try:
            data = f.read()
            rows = csv_rows(data)
            if not rows:
                continue
            hdr = [str(c).strip().lower() for c in rows[0]]
            ic_col = next((i for i, h in enumerate(hdr) if "old" in h and "i" in h), None)
            ac_col = next((i for i, h in enumerate(hdr) if "account" in h and "no" in h), None)
            if ic_col is None or ac_col is None:
                continue
            m = {}
            for r in rows[1:]:
                ch = _txt(r[ic_col] if ic_col < len(r) else "").strip().upper()
                ac = _txt(r[ac_col] if ac_col < len(r) else "").strip()
                if ch and ac and ac.lower() not in BLANKS:
                    m[ch] = ac
            save_acct_map(m)
            added += len(m)
        except Exception:
            pass
    ss.auto_dg_msg = ("success", f"Harvested {added:,} account mappings from CSV(s).")


# ---------------------------------------------------------------- UI
def render_autostat_tab():
    ss = st.session_state
    ss.setdefault("auto_fvr", None)
    ss.setdefault("auto_id", None)
    ss.setdefault("auto_msg", None)

    st.subheader("📋 EWB Field Result Autostat Builder")
    st.caption("Daily: drop the FIELD RESULT file → pick the date → download the autostat upload.")

    with st.sidebar:
        st.markdown("---")
        st.markdown("**⚙️ Autostat Settings**")
        ss["auto_bank"] = st.text_input("Bank filter", value=ss.get("auto_bank", AUTO_BANK), key="auto_bank_inp")
        ss["auto_by"] = st.text_input("REMARKS BY code", value=ss.get("auto_by", AUTO_REMARKS_BY), key="auto_by_inp")
        ss["file_password"] = st.text_input("File password", value=ss.get("file_password", DEFAULT_PASSWORD),
                                            type="password", key="auto_pw_inp")
        ss["auto_protect"] = st.checkbox("Password-protect the download", value=False, key="auto_protect_cb")

    import datetime as _dt
    day = st.date_input("Date to extract", value=_dt.date.today(), key="auto_date_tab")
    use_visit = (st.radio("Visit time in REMARKS DATE",
                          ["Use actual visit time from the field file",
                           "Fixed time (as in the template)"],
                          key="auto_time_mode_tab", horizontal=True) != "Fixed time (as in the template)")

    f = st.file_uploader("Drop the FIELD RESULT file (.xlsx) — the full monthly workbook is fine.",
                         type=["xlsx"], key="auto_up_tab")
    if f is not None:
        fid = (f.name, f.size, ss.get("auto_bank"))
        if ss.get("auto_id") != fid:
            try:
                with st.spinner("Reading the field result file..."):
                    ss.auto_fvr = read_fvr(f, ss.get("auto_bank", AUTO_BANK), ss.get("file_password", DEFAULT_PASSWORD))
                ss.auto_id = fid
                known = ss.auto_fvr[~ss.auto_fvr["ACCOUNT NUMBER"].str.lower().isin(BLANKS)]
                save_acct_map(dict(zip(known["chcode"], known["ACCOUNT NUMBER"])))
                ss.auto_msg = None
            except Exception as e:
                ss.auto_fvr, ss.auto_id = None, None
                ss.auto_msg = ("error", f"Could not read that file: {e}")

    st.file_uploader("Drop the DataGrid export (.xlsx) to fill missing account numbers",
                     type=["xlsx"], key="auto_dg_tab", on_change=dg_load)
    dgm = load_dg_map()
    st.caption(f"DataGrid on file: {len(dgm):,} accounts" + (
        "" if dgm else " — none yet. Drop a DataGrid above."))

    with st.expander("Optional: fill missing accounts from Daily Remark CSV"):
        st.file_uploader("Daily_Remark .csv", type=["csv"], accept_multiple_files=True,
                         key="auto_csv_tab", on_change=auto_harvest)

    dg_msg = ss.pop("auto_dg_msg", None)
    if dg_msg:
        getattr(st, dg_msg[0])(dg_msg[1])
    if ss.get("auto_msg"):
        getattr(st, ss.auto_msg[0])(ss.auto_msg[1])

    fvr = ss.get("auto_fvr")
    if fvr is None or fvr.empty:
        st.info("Drop the field result file to start." if fvr is None
                else f'No "{ss.get("auto_bank", AUTO_BANK)}" rows found in that file.')
        return

    status_map = load_status_map()
    ready, hold, stats = build_autostat(
        fvr, day, status_map, load_acct_map(),
        ss.get("auto_by", AUTO_REMARKS_BY), use_visit,
        ss.get("auto_fixed_time", AUTO_TIME) or AUTO_TIME,
        dg_map=dgm, use_dg=True)

    per_date = fvr.groupby("DATE").size()
    if stats["rows"] == 0:
        last = per_date.index.max()
        st.warning(f"No rows for {pd.Timestamp(day):%m/%d/%Y}. Latest in the file: "
                   f"{last:%m/%d/%Y} ({int(per_date.loc[last]):,} rows). Change the date above.")
        return

    m = st.columns(7)
    m[0].metric("Rows for date", f"{stats['rows']:,}")
    m[1].metric("CANCEL removed", f"{stats['cancel']:,}")
    m[2].metric("Duplicate refs removed", f"{stats['dup']:,}")
    m[3].metric("Visits retimed +1 min", f"{stats['retimed']:,}")
    m[4].metric("Ready for autostat", f"{stats['ready']:,}")
    m[5].metric("Status not in REF", f"{stats['hold_status']:,}")
    m[6].metric("Account no. missing", f"{stats['hold_acct']:,}")

    aligned = stats["aligned"]
    if aligned:
        n_al = sum(aligned.values())
        with st.expander(f"{n_al:,} row(s) auto-aligned to nearest STATUS CODE — click to check/change"):
            opts = sorted(set(status_map.values()))
            tbl = pd.DataFrame({"STATUS - SUB STATUS": list(aligned),
                                "Rows": list(aligned.values()),
                                "STATUS CODE": [status_map[k] for k in aligned]})
            ed0 = st.data_editor(tbl, hide_index=True, disabled=["STATUS - SUB STATUS", "Rows"],
                                 column_config={"STATUS CODE": st.column_config.SelectboxColumn(options=opts)},
                                 key="auto_aligned_" + str(abs(hash(tuple(aligned)))))
            if st.button("Save changes", key="auto_save_aligned_tab"):
                save_status_map(dict(zip(ed0["STATUS - SUB STATUS"], ed0["STATUS CODE"])))
                st.rerun()

    bad = hold[hold["REASON"].str.startswith("Status")]
    if not bad.empty:
        st.warning(f"{bad['KEY'].nunique()} status combination(s) not in REF — "
                   f"{stats['hold_status']:,} row(s) held back. Pick a STATUS CODE and Save.")
        options = [""] + sorted(set(status_map.values()))
        grp = bad.groupby("KEY").size().sort_values(ascending=False)
        table = pd.DataFrame({"STATUS - SUB STATUS": grp.index, "Rows": grp.values,
                              "STATUS CODE": [""] * len(grp)})
        ed = st.data_editor(table, hide_index=True, disabled=["STATUS - SUB STATUS", "Rows"],
                            column_config={"STATUS CODE": st.column_config.SelectboxColumn(options=options)},
                            key="auto_status_" + str(abs(hash(tuple(grp.index)))))
        if st.button("Save status codes", key="auto_save_status_tab"):
            save_status_map(dict(zip(ed["STATUS - SUB STATUS"], ed["STATUS CODE"])))
            st.rerun()

    miss = hold[hold["REASON"].str.contains("account", case=False)].drop_duplicates("CH CODE")
    if not miss.empty:
        st.warning(f"{len(miss):,} CH code(s) have no account number. Type them below and Save.")
        t2 = miss[["CH CODE", "CH NAME", "KEY"]].assign(**{"ACCOUNT NUMBER": ""})
        ed2 = st.data_editor(t2, hide_index=True, disabled=["CH CODE", "CH NAME", "KEY"],
                             key="auto_acct_" + str(abs(hash(tuple(miss["CH CODE"])))))
        if st.button("Save account numbers", key="auto_save_acct_tab"):
            save_acct_map(dict(zip(ed2["CH CODE"], ed2["ACCOUNT NUMBER"])))
            st.rerun()

    pw = ss.get("file_password", DEFAULT_PASSWORD) if ss.get("auto_protect", False) else ""
    d1, d2 = st.columns(2)
    d1.download_button(
        f"📥 Download Autostat ({stats['ready']:,} rows)",
        data=autostat_workbook(ready, pw) if stats["ready"] else b"",
        file_name=f"EWB_FIELD_RESULT_AUTOSTAT_{pd.Timestamp(day):%m-%d-%Y}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary", disabled=not stats["ready"],
        key="dl_autostat_tab")
    if not hold.empty:
        d2.download_button(
            f"📥 Download {len(hold):,} held-back rows (Needs Attention)",
            data=attention_workbook(hold),
            file_name=f"EWB_AUTOSTAT_NEEDS_ATTENTION_{pd.Timestamp(day):%m-%d-%Y}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="dl_attention_tab")

    if not ready.empty:
        st.dataframe(ready, hide_index=True, use_container_width=True)
