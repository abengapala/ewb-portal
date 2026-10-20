"""EWB BAU Daily Status Report Tab — for use inside the EWB Portal multi-page app.

All logic from Report DRR.py is preserved exactly.
set_page_config, st.title, and __main__ guard removed for safe import.
"""
import csv
import io
import re
import zipfile
from pathlib import Path
from datetime import datetime, time

import msoffcrypto
import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

DRR_SHEET = "PASTE DRR HERE"
REP_SHEET = "STATUS REP"
COLS = ["S.No", "Date", "Time", "Debtor", "Account No.", "Card No.",
        "Service No.", "DPD", "Call Status", "Status", "Remark"]
REP_COLS = ["DATE", "Account No.", "Account Name", "Remarks"]
AGENT = "SP MADRID"
DROP_DEBTOR = "EWB AUTO"
DEFAULT_PASSWORD = "spmadrid1234"
OPEN_PASSWORDS = [DEFAULT_PASSWORD, "SPM2026", "SPM 2026"]
DATA_DIR = Path(__file__).resolve().parent.parent / "ewb_data"
MASTER = DATA_DIR / "EWB_BAU_DAILY_STATUS_REPORT_master.xlsx"
REPORTS_DIR = DATA_DIR / "reports"
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_REMARK_PREFIX = re.compile(r"^\s*SP MADRID\s+\d{1,2}[-/]\d{1,2}[-/]\d{2,4}\s*", re.I)


# ---------------------------------------------------------------- parsing
def to_date(v, dayfirst=False):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return pd.NaT
    if isinstance(v, (datetime, pd.Timestamp)):
        return pd.Timestamp(v).normalize()
    s = str(v).strip()
    if not s:
        return pd.NaT
    if re.fullmatch(r"\d{5}(\.\d+)?", s):
        return (pd.Timestamp("1899-12-30") + pd.to_timedelta(float(s), unit="D")).normalize()
    if re.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}", s):
        try:
            return pd.Timestamp(s[:10]).normalize()
        except Exception:
            pass
    fmts = ("%d/%m/%Y", "%m/%d/%Y", "%d/%m/%y", "%m/%d/%y") if dayfirst else \
           ("%m/%d/%Y", "%d/%m/%Y", "%m/%d/%y", "%d/%m/%y")
    for f in fmts:
        try:
            return pd.Timestamp(datetime.strptime(s[:10].replace("-", "/"), f)).normalize()
        except ValueError:
            pass
    try:
        return pd.to_datetime(s).normalize()
    except Exception:
        return pd.NaT


def _txt(x):
    if isinstance(x, time):
        return x.strftime("%H:%M:%S")
    if x is None:
        return ""
    if isinstance(x, float) and x.is_integer():
        x = int(x)
    return _ILLEGAL.sub("", str(x).strip())


def _nonblank(rows):
    return [r for r in rows if any(str(c).strip() for c in r if c is not None)]


def make_remarks(dates, status):
    d = dates.dt.strftime("%m/%d/%Y").fillna("")
    return (f"{AGENT} " + d + " " + status).str.strip()


def detect_dayfirst(values, default):
    first = second = 0
    for v in values:
        m = re.match(r"^\s*(\d{1,2})[/-](\d{1,2})[/-]\d{2,4}", str(v))
        if m:
            first += int(m.group(1)) > 12
            second += int(m.group(2)) > 12
    if first and not second:
        return True
    if second and not first:
        return False
    return default


def drr_to_report(rows, dayfirst_default=False):
    rows = _nonblank(rows)
    if rows and any("account" in str(c).lower() for c in rows[0]):
        hdr = [str(c).strip().lower() for c in rows[0]]
        idx = [hdr.index(n.lower()) if n.lower() in hdr else i for i, n in enumerate(COLS)]
        fixed = [[(r[j] if j < len(r) else "") for j in idx] for r in rows[1:]]
    else:
        fixed = [(list(r) + [""] * len(COLS))[:len(COLS)] for r in rows]
    df = pd.DataFrame(fixed, columns=COLS)
    for c in COLS:
        if c != "Date":
            df[c] = df[c].map(_txt)
    dayfirst = detect_dayfirst(df["Date"].unique(), dayfirst_default)
    cache = {v: to_date(v, dayfirst) for v in df["Date"].unique()}
    date = pd.to_datetime(df["Date"].map(cache))
    return pd.DataFrame({
        "DATE": date,
        "Account No.": df["Account No."],
        "Account Name": df["Debtor"],
        "Remarks": make_remarks(date, df["Status"]),
    })


def report_to_frame(rows):
    rows = _nonblank(rows)
    if rows and str(rows[0][0]).strip().lower() == "date":
        rows = rows[1:]
    fixed = [(list(r) + [""] * 4)[:4] for r in rows]
    df = pd.DataFrame(fixed, columns=REP_COLS)
    for c in REP_COLS[1:]:
        df[c] = df[c].map(_txt)
    cache = {v: to_date(v) for v in df["DATE"].unique()}
    date = pd.to_datetime(df["DATE"].map(cache))
    status = df["Remarks"].str.replace(_REMARK_PREFIX, "", regex=True)
    return pd.DataFrame({"DATE": date, "Account No.": df["Account No."],
                         "Account Name": df["Account Name"],
                         "Remarks": make_remarks(date, status)})


def finalize(rep, drop_auto=True, digits_only=True, prefix="", dedup=True):
    stats = {"auto": 0, "chassis": 0, "dup": 0}
    rep = rep.copy()
    rep["Account No."] = rep["Account No."].str.replace(r"\s+", "", regex=True)
    if drop_auto:
        keep = ~rep["Account Name"].str.upper().str.contains(DROP_DEBTOR, regex=False)
        stats["auto"] = int((~keep).sum())
        rep = rep[keep]
    if digits_only or prefix:
        acct = rep["Account No."]
        keep = acct.str.fullmatch(r"\d*[1-9]\d*") if digits_only else acct != ""
        if prefix:
            keep &= acct.str.startswith(prefix)
        stats["chassis"] = int((~keep).sum())
        rep = rep[keep]
    if dedup:
        before = len(rep)
        rep = rep.drop_duplicates(subset=["DATE", "Account No.", "Remarks"], keep="first")
        stats["dup"] = before - len(rep)
    return rep.reset_index(drop=True), stats


def parse_pasted(text, **kw):
    rows = list(csv.reader(io.StringIO(text.replace("\r\n", "\n")), delimiter="\t"))
    return finalize(drr_to_report(rows), **kw)


def _bytes(file):
    if isinstance(file, (str, Path)):
        return Path(file).read_bytes()
    if hasattr(file, "seek"):
        file.seek(0)
    return file.read()


def csv_rows(data):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    return list(csv.reader(io.StringIO(text)))


def open_file_wb(file, password=DEFAULT_PASSWORD):
    data = _bytes(file)
    bio = io.BytesIO(data)
    if not zipfile.is_zipfile(bio):
        bio.seek(0)
        ms = msoffcrypto.OfficeFile(bio)
        if not ms.is_encrypted():
            raise ValueError("This is not a valid .xlsx file.")
        last = None
        for pw in dict.fromkeys([password, *OPEN_PASSWORDS]):
            if not pw:
                continue
            try:
                bio.seek(0)
                ms = msoffcrypto.OfficeFile(bio)
                ms.load_key(password=pw)
                dec = io.BytesIO()
                ms.decrypt(dec)
                dec.seek(0)
                return load_workbook(dec, data_only=True, read_only=True)
            except msoffcrypto.exceptions.InvalidKeyError as e:
                last = e
        raise ValueError("Wrong password. Type the correct one in the sidebar.") from last
    bio.seek(0)
    return load_workbook(bio, data_only=True, read_only=True)


def read_workbook(file, password=DEFAULT_PASSWORD, **kw):
    data = _bytes(file)
    name = str(getattr(file, "name", file)).lower()
    if name.endswith(".csv") or not (data[:2] == b"PK" or data[:4] == b"\xd0\xcf\x11\xe0"):
        out, stats = finalize(drr_to_report(csv_rows(data), dayfirst_default=True), **kw)
        stats["baddate"] = 0
        return out, stats
    wb = open_file_wb(io.BytesIO(data), password)
    rep_ws = None
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, min_col=1, max_col=4, values_only=True), None)
        if first and [str(c or "").strip().lower() for c in first] == [c.lower() for c in REP_COLS]:
            rep_ws = ws
            break
    baddate = 0
    if rep_ws is not None:
        rep = report_to_frame(list(rep_ws.iter_rows(min_col=1, max_col=4, values_only=True)))
        bad = rep["DATE"].isna()
        baddate = int(bad.sum())
        rep = rep[~bad]
        out, stats = finalize(rep, **kw)
        stats["baddate"] = baddate
        return out, stats
    # DRR xlsx
    drr_ws = next((ws for ws in wb.worksheets if DRR_SHEET.lower() in ws.title.lower()), None) or wb.worksheets[0]
    rows = list(drr_ws.iter_rows(values_only=True))
    out, stats = finalize(drr_to_report(rows), **kw)
    stats["baddate"] = baddate
    return out, stats


def to_excel(df, password=""):
    wb = Workbook()
    ws = wb.active
    ws.title = REP_SHEET
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")
    ws.append(REP_COLS)
    for cell in ws[1]:
        cell.fill = hdr_fill
        cell.font = hdr_font
    for _, row in df.iterrows():
        date_str = row["DATE"].strftime("%m/%d/%Y") if pd.notna(row["DATE"]) else ""
        ws.append([date_str, row["Account No."], row["Account Name"], row["Remarks"]])
    buf = io.BytesIO()
    wb.save(buf)
    if not password:
        return buf.getvalue()
    buf.seek(0)
    out = io.BytesIO()
    msoffcrypto.OfficeFile(buf).encrypt(password, out)
    return out.getvalue()


def report_name(df):
    return f"EWB_BAU_DAILY_STATUS_REPORT_as_of_{df['DATE'].max():%B_%d__%Y}.xlsx"


def save_report(df, password=DEFAULT_PASSWORD):
    data = to_excel(df, password or DEFAULT_PASSWORD)
    DATA_DIR.mkdir(exist_ok=True)
    REPORTS_DIR.mkdir(exist_ok=True)
    MASTER.write_bytes(data)
    (REPORTS_DIR / report_name(df)).write_bytes(data)


# ---------------------------------------------------------------- UI
def render_drr_tab():
    ss = st.session_state
    ss.setdefault("drr_df", pd.DataFrame(columns=REP_COLS))
    ss.setdefault("drr_msg", None)
    ss.setdefault("drr_loaded_id", None)
    ss.setdefault("drr_paste", "")
    ss.setdefault("drr_up_key", 0)
    ss.setdefault("drr_replace_ok", True)
    ss.setdefault("drr_master_checked", False)

    def filt():
        return dict(drop_auto=ss.get("drr_drop_auto", True),
                    digits_only=ss.get("drr_digits_only", True),
                    prefix=ss.get("drr_prefix", "").strip(),
                    dedup=ss.get("drr_dedup", True))

    def removed_text(stats):
        return (f"removed {stats['auto']:,} {DROP_DEBTOR} row(s), {stats['chassis']:,} chassis/invalid "
                f"account row(s) and {stats['dup']:,} repeated effort(s)")

    if not ss.drr_master_checked:
        ss.drr_master_checked = True
        if MASTER.exists():
            try:
                ss.drr_df, _ = read_workbook(MASTER, DEFAULT_PASSWORD, **filt())
                ss.drr_msg = ("info", f"Loaded saved report: {len(ss.drr_df):,} rows up to "
                              f"{ss.drr_df['DATE'].max():%m/%d/%Y}. Drop today's DRR to add it.")
            except Exception as e:
                ss.drr_msg = ("error", f"Could not load saved report: {e}")

    st.subheader("📊 EWB BAU Daily Status Report")
    st.caption("Upload or paste today's MC3 DRR — cleaned, merged, and ready to download.")

    with st.sidebar:
        st.markdown("---")
        st.markdown("**⚙️ BAU Report Settings**")
        st.checkbox(f'Delete rows where debtor is "{DROP_DEBTOR}"', value=True, key="drr_drop_auto")
        st.checkbox("Delete chassis numbers (Account No. not all digits)", value=True, key="drr_digits_only")
        st.text_input("Only keep accounts starting with (e.g. 200)", value="", key="drr_prefix")
        st.checkbox("Keep same effort only once per account per day", value=True, key="drr_dedup")
        st.text_input("File password (protects download & opens protected files)",
                      value=DEFAULT_PASSWORD, type="password", key="drr_file_password")
        st.checkbox("Password-protect the downloaded file", value=True, key="drr_protect_out")
        up = st.file_uploader("Set/replace saved report (e.g. Sept 30 file)", type=["xlsx"], key="drr_rep_up")
        if up is not None:
            fid = (up.name, up.size)
            if ss.drr_loaded_id != fid:
                try:
                    ss.drr_df, stats = read_workbook(up, ss.get("drr_file_password", DEFAULT_PASSWORD), **filt())
                    ss.drr_loaded_id = fid
                    save_report(ss.drr_df, ss.get("drr_file_password", DEFAULT_PASSWORD))
                    ss.drr_msg = ("success", f"Loaded {len(ss.drr_df):,} rows from {up.name}; "
                                  f"{removed_text(stats)}. Saved — just drop the DRR from now on.")
                except Exception as e:
                    ss.drr_msg = ("error", f"Could not read that file: {e}")
        if st.button("Start a new report (clear everything)", key="drr_clear_btn"):
            if MASTER.exists():
                MASTER.rename(MASTER.with_name(f"master_backup_{datetime.now():%Y%m%d_%H%M%S}.xlsx"))
            ss.drr_df = pd.DataFrame(columns=REP_COLS)
            ss.drr_loaded_id = None
            ss.drr_msg = ("info", "Cleared. Starting a new report.")
            st.rerun()

    def merge_new(new, stats):
        if new.empty:
            ss.drr_msg = ("warning", f"No rows left to add ({removed_text(stats)}).")
            return False
        bad = int(new["DATE"].isna().sum())
        if bad:
            ss.drr_msg = ("error", f"{bad:,} row(s) have an unreadable date.")
            return False
        overlap = sorted(set(new["DATE"]) & set(ss.drr_df["DATE"]))
        if overlap and not ss.get("drr_replace_ok"):
            ss.drr_msg = ("warning", f"The report already has rows for {len(overlap)} date(s). "
                          "Tick 'Replace existing dates' below, then press Add again.")
            return False
        old = ss.drr_df[~ss.drr_df["DATE"].isin(overlap)]
        merged = pd.concat([old, new], ignore_index=True)
        merged["DATE"] = pd.to_datetime(merged["DATE"])
        ss.drr_df = merged.sort_values("DATE", kind="stable", ignore_index=True)
        days = sorted(set(new["DATE"]))
        span = f"{days[0]:%m/%d/%Y}" + (f" to {days[-1]:%m/%d/%Y}" if len(days) > 1 else "")
        ss.drr_msg = ("success", f"Added {len(new):,} rows ({span}); {removed_text(stats)}. "
                      f"Total: {len(ss.drr_df):,} rows. Saved.")
        try:
            save_report(ss.drr_df, ss.get("drr_file_password", DEFAULT_PASSWORD))
        except Exception as e:
            ss.drr_msg = ("warning", ss.drr_msg[1].replace("Saved.", f"BUT could not save: {e}"))
        return True

    def add_pasted():
        if not ss.drr_paste.strip():
            ss.drr_msg = ("warning", "Nothing pasted.")
            return
        new, stats = parse_pasted(ss.drr_paste, **filt())
        if merge_new(new, stats):
            ss.drr_paste = ""

    def add_uploaded():
        files = ss.get(f"drr_day_up_{ss.drr_up_key}") or []
        if not files:
            return
        notes, level, all_ok = [], "success", True
        for f in files:
            try:
                new, stats = read_workbook(f, ss.get("drr_file_password", DEFAULT_PASSWORD), **filt())
            except Exception as e:
                notes.append(f"{f.name}: could not read ({e})")
                level, all_ok = "error", False
                continue
            ok = merge_new(new, stats)
            lvl, text = ss.drr_msg
            notes.append(f"{f.name}: {text}")
            if not ok:
                all_ok = False
                level = "error" if lvl == "error" or level == "error" else "warning"
        ss.drr_msg = (level, "  \n".join(notes))
        if all_ok:
            ss.drr_up_key += 1

    tab_up, tab_paste = st.tabs(["📂 Upload DRR file (recommended)", "📋 Paste DRR from Excel"])
    with tab_up:
        st.file_uploader(
            "Drop today's MC3 DRR here (Daily_Remark .csv or .xlsx) — auto-cleaned and added",
            type=["csv", "xlsx"], accept_multiple_files=True,
            key=f"drr_day_up_{ss.drr_up_key}", on_change=add_uploaded)
    with tab_paste:
        st.text_area("Paste DRR rows copied from Excel (header row optional; columns A:K only)",
                     key="drr_paste", height=180)
        st.button("Add to report", type="primary", on_click=add_pasted, key="drr_add_paste")

    st.checkbox("Replace existing dates when re-uploading today's DRR (safe to reuse)", key="drr_replace_ok")

    if ss.drr_msg:
        getattr(st, ss.drr_msg[0])(ss.drr_msg[1])

    df = ss.drr_df
    if df.empty:
        st.info("No data yet. Add a DRR above, or upload last month's report in the sidebar.")
    else:
        pw = ss.get("drr_file_password", DEFAULT_PASSWORD) if ss.get("drr_protect_out", True) else ""
        st.download_button(
            "📥 Download status report",
            data=to_excel(df, pw),
            file_name=report_name(df),
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            key="drr_download_btn"
        )
        st.caption(f"Auto-saved in: {DATA_DIR}")
        per_day = (df.groupby(df["DATE"].dt.strftime("%m/%d/%Y"), sort=False).size()
                   .rename("Rows").reset_index().rename(columns={"DATE": "Date"}))
        a, b = st.columns([1, 3])
        a.subheader("Days in report")
        a.dataframe(per_day, hide_index=True)
        b.subheader("Latest 200 rows")
        b.dataframe(df.tail(200).assign(DATE=lambda x: x["DATE"].dt.strftime("%m/%d/%Y")), hide_index=True)
